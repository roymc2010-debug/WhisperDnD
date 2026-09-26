"""Unit tests for Word docx chronicle exporter."""

import os
import shutil
import tempfile
import unittest
from pathlib import Path
import docx

from src.exporters.docx_exporter import export_chronicle_docx


class TestDocxExporter(unittest.TestCase):
    def setUp(self):
        self.temp_dir = tempfile.mkdtemp()

    def tearDown(self):
        shutil.rmtree(self.temp_dir, ignore_errors=True)

    def test_export_chronicle_docx(self):
        sample_md = """# Crónica de Prueba
## 1. Resumen Narrativo
Los aventureros exploraron las cavernas olvidadas.
- Encontraron un cofre cerrado con runas.
- Derrotaron a tres goblins centinelas.
"""
        sample_roster = [
            {"player_name": "Rodrigo", "role": "DM", "character_name": "(DM)", "species": "N/A", "subclass": "N/A"},
            {"player_name": "Sofía", "role": "Pícaro", "character_name": "Kaelen", "species": "Elfo", "subclass": "Phantom"},
        ]

        test_out = Path(self.temp_dir) / "test_export.docx"
        try:
            exported_path = export_chronicle_docx(
                chronicle_md=sample_md,
                roster=sample_roster,
                session_date="19/09/2026",
                output_path=str(test_out),
            )

            self.assertTrue(Path(exported_path).is_file())
            # Read document back to verify contents
            doc = docx.Document(exported_path)
            full_text = "\n".join([p.text for p in doc.paragraphs])
            self.assertIn("Crónica de Sesión", full_text)
            self.assertIn("Resumen Narrativo", full_text)
            self.assertIn("Los aventureros exploraron", full_text)

            # Check table
            self.assertEqual(len(doc.tables), 1)
            self.assertEqual(len(doc.tables[0].rows), 3)  # header + 2 members
            self.assertEqual(len(doc.tables[0].columns), 5)  # 5 columns (Jugador, Personaje, Especie, Clase, Subclase)
        finally:
            if test_out.is_file():
                test_out.unlink()

    def test_export_3_tier_chronicle_docx(self):
        sample_3tier_md = """# SECCIÓN 1: CRÓNICA DETALLADA DE LA AVENTURA
## 1.1 Resumen Narrativo
El combate en las ruinas fue feroz.
## 1.2 Desglose de Combates
Markus acertó un golpe crítico con su mandoble.
# SECCIÓN 2: "PREVIOUSLY ON..." (RECAP PARA LA PRÓXIMA SESIÓN)
Los aventureros se adentraron en las sombras del bosque.
# SECCIÓN 3: RETROALIMENTACIÓN PERSONAL PARA MARKUS VEYL (RODRIGO)
## 3.1 Aciertos
Excelente uso de Riposte.
## 3.2 Áreas de Mejora
Recuerda utilizar Action Surge.
"""
        sample_roster = [
            {"player_name": "Roymc89", "role": "Guerrero", "character_name": "Markus Veyl", "species": "Humano", "subclass": "Battle Master"}
        ]
        test_out = Path(self.temp_dir) / "test_3tier_export.docx"
        try:
            exported_path = export_chronicle_docx(
                chronicle_md=sample_3tier_md,
                roster=sample_roster,
                session_date="19/09/2026",
                output_path=str(test_out),
            )
            self.assertTrue(Path(exported_path).is_file())
            doc = docx.Document(exported_path)
            full_text = "\n".join([p.text for p in doc.paragraphs])
            self.assertIn("SECCIÓN 1: CRÓNICA DETALLADA", full_text)
            self.assertIn("SECCIÓN 2: \"PREVIOUSLY ON...\"", full_text)
            self.assertIn("SECCIÓN 3: RETROALIMENTACIÓN PERSONAL", full_text)
            self.assertIn("MARKUS VEYL", full_text)
        finally:
            if test_out.is_file():
                test_out.unlink()

    def test_export_academic_notes_docx(self):
        from src.exporters.docx_exporter import export_academic_notes_docx

        sample_notes_md = """# 🎓 GUÍA DE ESTUDIO UNIVERSITARIA
## 1. Resumen de la Clase
Se introdujeron los diagramas de Bode y el margen de fase.
## 2. Conceptos Teóricos Fundamentales
- Estabilidad de Nyquist.
- Frecuencia de corte.
## 3. Fórmulas, Ecuaciones y Procedimientos
$$G(s) = \\frac{K}{s(s+1)}$$
## 4. Ejemplos Resueltos en Clase
Cálculo analítico del margen de ganancia.
## 5. Avisos Relevantes y Fechas Clave
Entrega de la práctica el 25/09.
## 6. Guía de Estudio Rápida
Revisar conceptos de polos complejos.
"""
        test_out = Path(self.temp_dir) / "test_academic_export.docx"
        try:
            exported_path = export_academic_notes_docx(
                notes_md=sample_notes_md,
                subject="Sistemas de Control",
                topic="Diagramas de Bode",
                lecture_date="19/09/2026",
                output_path=str(test_out),
            )
            self.assertTrue(Path(exported_path).is_file())
            doc = docx.Document(exported_path)
            full_text = "\n".join([p.text for p in doc.paragraphs])
            self.assertIn("Sistemas de Control", full_text)
            self.assertIn("Diagramas de Bode", full_text)
            self.assertIn("Resumen de la Clase", full_text)
            self.assertIn("Conceptos Teóricos", full_text)
            self.assertIn("Avisos Relevantes", full_text)
        finally:
            if test_out.is_file():
                test_out.unlink()

    def test_export_living_journal_docx_deduplicates_chapters(self):
        from src.exporters.docx_exporter import export_living_journal_docx

        # Campaign state with duplicate session entries for Session 1
        campaign_state = {
            "campaign_name": "Campaña Deduplicada",
            "last_session": 1,
            "universal_pcs": [
                {
                    "personaje": "Halendiel Fang",
                    "jugador": "Liam O'Brien",
                    "especie": "Medio Elfo",
                    "clase": "Bardo",
                    "subclase": "Elocuencia",
                    "debut_sesion": 1,
                    "hitos_acumulados": ["Sesión #1: Inspiró a la compañía"],
                }
            ],
            "quests": [],
            "npcs": [],
            "sessions": [
                {
                    "session_number": 1,
                    "title": "Versión Vieja",
                    "date": "18/09/2026",
                    "chronicle_text": "Crónica anterior que debe ser reemplazada.",
                },
                {
                    "session_number": 1,
                    "title": "Versión Definitiva",
                    "date": "19/09/2026",
                    "chronicle_text": "Crónica nueva y única para el capítulo 1.",
                },
            ],
        }

        test_out = Path(self.temp_dir) / "test_dedup_living_journal.docx"
        try:
            exported_path = export_living_journal_docx(campaign_state, output_path=str(test_out))
            self.assertTrue(Path(exported_path).is_file())
            doc = docx.Document(exported_path)
            full_text = "\n".join([p.text for p in doc.paragraphs])
            
            # Exactly one Chapter 1 heading
            self.assertIn("Capítulo 1: Versión Definitiva", full_text)
            self.assertNotIn("Versión Vieja", full_text)
            # Count occurrences of "Capítulo 1:"
            ch1_count = sum(1 for p in doc.paragraphs if "Capítulo 1:" in p.text)
            self.assertEqual(ch1_count, 1)
        finally:
            if test_out.is_file():
                test_out.unlink()

    def test_clean_math_delimiters_for_docx(self):
        from src.exporters.docx_exporter import clean_math_delimiters_for_docx
        raw = r"La ganancia esperada es $$\mathbb{E}[X] = \sum x_i \cdot p_i = \$1,000$$ con valor \(P(A \le 5)\)."
        cleaned = clean_math_delimiters_for_docx(raw)
        self.assertNotIn("$$", cleaned)
        self.assertNotIn(r"\(", cleaned)
        self.assertNotIn(r"\)", cleaned)
        self.assertIn("$1,000", cleaned)
        self.assertIn("≤", cleaned)

    def test_export_code_block_consolas_formatting(self):
        from src.exporters.docx_exporter import export_academic_notes_docx
        from docx.shared import Pt

        sample_code_md = """# Guía de Algoritmos
## 1. Implementación de Árbol
```python
def dfs(node):
    if not node:
        return
    print(node.val)
```
Fin de la explicación.
"""
        test_out = Path(self.temp_dir) / "test_code_export.docx"
        try:
            exported_path = export_academic_notes_docx(
                notes_md=sample_code_md,
                subject="Algoritmos",
                topic="DFS",
                output_path=str(test_out),
            )
            doc = docx.Document(exported_path)
            # Find code paragraphs
            code_paras = [p for p in doc.paragraphs if "def dfs" in p.text or "print(node.val)" in p.text]
            self.assertTrue(len(code_paras) >= 2)
            for p in code_paras:
                self.assertEqual(p.paragraph_format.line_spacing, 1.0)
                self.assertEqual(p.paragraph_format.space_after, Pt(0))
                self.assertTrue(len(p.runs) >= 1)
                self.assertEqual(p.runs[0].font.name, "Consolas")
                self.assertEqual(p.runs[0].font.size, Pt(8.5))
        finally:
            if test_out.is_file():
                test_out.unlink()

    def test_export_markdown_table_formatting(self):
        from src.exporters.docx_exporter import export_academic_notes_docx

        sample_table_md = """# Guía de Modelos
## 1. Comparativa de Modelos
| Modelo | Parámetros | Precisión |
|---|---|---|
| GPT-4 | 1.8T | 92.5% |
| Claude 3.5 | 500B | 93.1% |

Conclusión comparativa.
"""
        test_out = Path(self.temp_dir) / "test_table_export.docx"
        try:
            exported_path = export_academic_notes_docx(
                notes_md=sample_table_md,
                subject="Inteligencia Artificial",
                topic="Modelos de Lenguaje",
                output_path=str(test_out),
            )
            doc = docx.Document(exported_path)
            # Should have created a table from the Markdown table
            self.assertTrue(len(doc.tables) >= 1)
            tbl = doc.tables[0]
            self.assertEqual(len(tbl.rows), 3)  # header + 2 rows
            self.assertEqual(len(tbl.columns), 3)
            self.assertEqual(tbl.rows[0].cells[0].text, "Modelo")
            self.assertEqual(tbl.rows[0].cells[1].text, "Parámetros")
            self.assertEqual(tbl.rows[1].cells[0].text, "GPT-4")
            self.assertEqual(tbl.rows[2].cells[0].text, "Claude 3.5")
            self.assertTrue(tbl.rows[0].cells[0].paragraphs[0].runs[0].font.bold)
        finally:
            if test_out.is_file():
                test_out.unlink()


if __name__ == "__main__":
    unittest.main()
